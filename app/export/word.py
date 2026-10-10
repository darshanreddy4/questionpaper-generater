"""DOCX export (python-docx). Mathematics is embedded as high-resolution images whose alt-text keeps the LaTeX."""
from __future__ import annotations

import io
import json
import re
from typing import Dict, List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app import languages
from app.diagrams import DiagramError, render_png
from app.exam.schema import Answer, PaperContent, Question, SubQuestion, fmt_marks
from app.export.common import duration_text, number_questions, option_letter, section_heading, section_marks_line
from app.export import indic
from app.export.labels import ENGLISH, Labels, for_language
from app.export.mathtext import latex_to_text, render_math, split_math

TEXT_WIDTH_MM = 174
GREEN = RGBColor(0x14, 0x53, 0x2D)
GREY = RGBColor(0x55, 0x55, 0x55)
BLUE = RGBColor(0x1F, 0x3A, 0x68)
_BOLD = re.compile(r"(\*\*)")
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class _Builder:
    def __init__(self, size: float, lab: Labels = ENGLISH):
        self.size = size
        self.lab = lab
        self.doc = Document()
        self._diagram_cache: Dict[str, Optional[object]] = {}
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(18)
        sec.top_margin, sec.bottom_margin = Mm(15), Mm(18)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        normal.font.size = Pt(size)
        normal.paragraph_format.space_after = Pt(2)
        normal.paragraph_format.space_before = Pt(0)
        indic.prepare_docx_styles(self.doc, size)

    # ---- text helpers

    def rich(self, par, text: str, bold: bool = False, italic: bool = False, color: Optional[RGBColor] = None,
             size: Optional[float] = None) -> None:
        for kind, content in split_math(text or ""):
            if kind == "text":
                is_bold = bold
                for piece in _BOLD.split(_CTRL.sub("", content)):
                    if piece == "**":
                        is_bold = not is_bold
                        continue
                    lines = piece.split("\n")
                    for i, line in enumerate(lines):
                        if i:
                            par.add_run().add_break()
                        if line:
                            run = par.add_run(line)
                            run.bold, run.italic = is_bold, italic
                            if color is not None:
                                run.font.color.rgb = color
                            if size:
                                run.font.size = Pt(size)
            else:
                img = render_math(content, size or self.size)
                if img is None:
                    run = par.add_run(latex_to_text(content))
                    run.italic = True
                    continue
                run = par.add_run()
                shape = run.add_picture(io.BytesIO(img.png), height=Pt(img.height_pt))
                shape._inline.docPr.set("descr", content[:250])
                rpr = run._r.get_or_add_rPr()
                pos = OxmlElement("w:position")
                pos.set(qn("w:val"), str(-int(round(img.drop_pt * 2))))
                rpr.append(pos)

    def para(self, text: str = "", *, align=None, left: float = 0, hanging: float = 0, bold=False, italic=False,
             color=None, size=None, space_before: float = 0, space_after: float = 2, keep_next=False):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        if align is not None:
            p.alignment = align
        pf.left_indent = Mm(left) if left else None
        pf.first_line_indent = Mm(-hanging) if hanging else None
        pf.space_before, pf.space_after = Pt(space_before), Pt(space_after)
        pf.keep_with_next = keep_next
        if text:
            self.rich(p, text, bold=bold, italic=italic, color=color, size=size)
        return p

    def shade(self, par, fill: str) -> None:
        ppr = par._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        ppr.append(shd)

    def rule(self) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        ppr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "111111")):
            bottom.set(qn(f"w:{k}"), v)
        borders.append(bottom)
        ppr.append(borders)

    def marks_tab(self, par, marks: float) -> None:
        par.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
        run = par.add_run(f"\t[{fmt_marks(marks)}]")
        run.bold = True

    # ---- blocks

    def diagram(self, spec: Optional[dict], left: float = 8) -> None:
        if not spec:
            return
        key = json.dumps(spec, sort_keys=True)
        if key not in self._diagram_cache:
            try:
                self._diagram_cache[key] = render_png(spec, dpi=220)
            except DiagramError:
                self._diagram_cache[key] = None
        rendered = self._diagram_cache[key]
        if rendered is None:
            return
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2, keep_next=bool(spec.get("caption")))
        width = min(rendered.width_in * 25.4, 110)
        p.add_run().add_picture(io.BytesIO(rendered.data), width=Mm(width))
        if spec.get("caption"):
            self.para(str(spec["caption"]), align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, color=GREY, size=9.5)

    def table(self, headers: List[str], rows: List[List[str]], shade: Optional[str] = "EEF1F6") -> None:
        if not rows:
            return
        ncol = max([len(headers)] + [len(r) for r in rows])
        t = self.doc.add_table(rows=0, cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if headers:
            cells = t.add_row().cells
            for i, h in enumerate(headers):
                cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
                self.rich(cells[i].paragraphs[0], h, bold=True, size=self.size - 1)
                tcpr = cells[i]._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), shade or "FFFFFF")
                tcpr.append(shd)
        for r in rows:
            cells = t.add_row().cells
            for i in range(ncol):
                cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
                self.rich(cells[i].paragraphs[0], r[i] if i < len(r) else "", size=self.size - 1)
        self.para(space_after=3)

    def options(self, options: List[str], correct: Optional[int], show_key: bool, left: float = 8) -> None:
        short = all(len(re.sub(r"\$.*?\$", "xx", o)) <= 26 for o in options)
        if short and len(options) % 2 == 0:
            for i in range(0, len(options), 2):
                p = self.para(left=left, space_after=1)
                p.paragraph_format.tab_stops.add_tab_stop(Mm(left + 70))
                for j in (i, i + 1):
                    if j > i:
                        p.add_run("\t")
                    self.rich(p, f"({option_letter(j)})\u2002{options[j]}", bold=show_key and correct == j)
                    if show_key and correct == j:
                        p.add_run(" \u2713").font.color.rgb = GREEN
        else:
            for j, opt in enumerate(options):
                p = self.para(left=left + 5, hanging=5, space_after=1)
                self.rich(p, f"({option_letter(j)})\u2002{opt}", bold=show_key and correct == j)
                if show_key and correct == j:
                    p.add_run(" \u2713").font.color.rgb = GREEN

    def answer(self, ans: Answer, options: List[str], total: float, left: float = 8) -> None:
        self.para(self.lab.get("answer"), bold=True, color=GREEN, size=8.5, left=left, space_before=2, space_after=0, keep_next=True)
        if ans.correct_option is not None and 0 <= ans.correct_option < len(options):
            self.para(f"{self.lab.get('correct_option')}: ({option_letter(ans.correct_option)}) {options[ans.correct_option]}",
                      color=GREEN, left=left, size=self.size - 0.5)
            if ans.text and not re.match(r"^\([a-dA-D]\)", ans.text):
                self.para(ans.text, color=GREEN, left=left, size=self.size - 0.5)
        elif ans.text:
            self.para(ans.text, color=GREEN, left=left, size=self.size - 0.5)
        steps = [s for s in ans.steps if s.text]
        if steps and not (len(steps) == 1 and steps[0].text.strip() == ans.text.strip()):
            rows = [[s.text, fmt_marks(s.marks)] for s in steps] + [[f"**{self.lab.get('total')}**", f"**{fmt_marks(total)}**"]]
            self.table([self.lab.get("marking_scheme"), self.lab.get("marks")], rows, shade="E9F3EC")
        self.diagram(ans.diagram)

    def subs(self, subs: List[SubQuestion], show_answers: bool) -> None:
        for s in subs:
            p = self.para(f"{s.label}\u2002{s.text}", left=14, hanging=6)
            self.marks_tab(p, s.marks)
            if s.options:
                self.options(s.options, s.answer.correct_option, show_answers, left=14)
            if show_answers:
                self.answer(s.answer, s.options, s.marks, left=14)

    def question_body(self, obj, number: Optional[int], marks: Optional[float], show_answers: bool) -> None:
        text = (f"**{number}.**\u2002" if number is not None else "") + obj.text
        p = self.para(text, left=8, hanging=8 if number is not None else 0, keep_next=bool(obj.options or obj.diagram))
        if marks is not None:
            self.marks_tab(p, marks)
        self.diagram(obj.diagram)
        if obj.table:
            self.table(obj.table.headers, obj.table.rows)
        match = getattr(obj, "match", None)
        if match:
            rows = [[f"{i + 1}. {a}", f"({option_letter(i)}) {b}"] for i, (a, b) in enumerate(zip(match.left, match.right))]
            self.table([self.lab.get("column_i"), self.lab.get("column_ii")], rows)
        if obj.options:
            self.options(obj.options, obj.answer.correct_option, show_answers)
        self.subs(obj.sub_questions, show_answers)
        if show_answers and not obj.sub_questions:
            total = getattr(obj, "marks", None) or sum(s.marks for s in obj.answer.steps)
            self.answer(obj.answer, obj.options, total)

    def footer(self, left_text: str, right_text: str) -> None:
        footer = self.doc.sections[0].footer
        p = footer.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM / 2), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
        page_fmt = self.lab.get("page", p="\2", t="\3")
        lead, rest = page_fmt.split("\2")
        mid, tail = rest.split("\3")
        r = p.add_run(left_text[:80] + "\t" + lead)
        r.font.size = Pt(8)
        for instr in ("PAGE", None, "NUMPAGES"):
            run = p.add_run()
            run.font.size = Pt(8)
            if instr is None:
                run.text = mid
                continue
            for kind, text in (("begin", None), (None, instr), ("end", None)):
                if kind:
                    el = OxmlElement("w:fldChar")
                    el.set(qn("w:fldCharType"), kind)
                else:
                    el = OxmlElement("w:instrText")
                    el.set(qn("xml:space"), "preserve")
                    el.text = text
                run._r.append(el)
        r = p.add_run(tail + "\t" + right_text)
        r.font.size = Pt(8)


def build_docx(content: PaperContent, mode: str = "question", font_size: float = 11.0) -> bytes:
    lab = for_language(content.meta.language)
    b = _Builder(font_size, lab)
    m = content.meta
    numbers = number_questions(content)
    center = WD_ALIGN_PARAGRAPH.CENTER

    if m.institution:
        b.para(m.institution, align=center, bold=True, color=BLUE, size=16, space_after=2)
    b.para(m.exam_name or m.title, align=center, bold=True, size=13, space_after=1)
    sub = " · ".join(x for x in (m.subject, m.class_level and f"{lab.get('class_level')}: {m.class_level}", m.board) if x)
    if sub:
        b.para(sub, align=center, color=GREY, size=10.5, space_after=1)
    if m.set_label:
        b.para(f"**{lab.get('set')} {m.set_label}**", align=center, size=10.5)
    if mode == "answer":
        b.para(lab.get("answer_key"), align=center, bold=True, color=GREEN, size=12, space_before=3)
    elif mode == "combined":
        b.para(lab.get("teachers_copy"), align=center, bold=True, color=GREEN, size=10.5, space_before=3)
    b.rule()
    p = b.para(f"**{lab.get('time_allowed')}:** {duration_text(m.duration_minutes, lab)}", bold=False, space_after=2)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(f"\t{lab.get('max_marks')}: {fmt_marks(m.total_marks)}").bold = True
    if mode == "question":
        p = b.para(f"{lab.get('name')}: ______________________________", space_before=4)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(95))
        p.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(f"\t{lab.get('roll_no')}: ____________")
        p.add_run("\t" + (f"{lab.get('date')}: {m.date}" if m.date else f"{lab.get('date')}: ________"))
    b.rule()
    if mode == "question" and m.instructions:
        b.para(lab.get("general_instructions"), bold=True, size=10, space_before=3, keep_next=True)
        for i, line in enumerate(m.instructions, 1):
            b.para(f"{i}.\u2002{line}", left=6, hanging=6, size=10.5, space_after=1)

    if mode == "answer":
        keys = [f"{numbers[q.id]}. ({option_letter(q.answer.correct_option)})"
                for s in content.sections for q in s.questions
                if q.type in ("mcq", "assertion_reason") and q.answer.correct_option is not None]
        if keys:
            b.para(lab.get("objective_key"), bold=True, size=10, space_before=6)
            cols = 8
            t = b.doc.add_table(rows=0, cols=cols)
            t.style = "Table Grid"
            for i in range(0, len(keys), cols):
                cells = t.add_row().cells
                for j in range(cols):
                    cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
                    if i + j < len(keys):
                        b.rich(cells[j].paragraphs[0], keys[i + j], size=b.size - 1)
            b.para(space_after=4)

    for sec in content.sections:
        if not sec.questions:
            continue
        head = b.para(section_heading(sec, lab) + (f" — {sec.title}" if sec.title else ""), align=center, bold=True,
                      color=BLUE, size=11, space_before=8, space_after=0, keep_next=True)
        b.shade(head, "EEF1F6")
        sub = b.para(section_marks_line(sec, lab), align=center, italic=True, color=GREY, size=10, space_after=3, keep_next=True)
        b.shade(sub, "EEF1F6")
        if sec.instruction and mode != "answer":
            b.para(sec.instruction, italic=True, left=8, size=10.5, space_after=3)
        for q in sec.questions:
            n = numbers[q.id]
            if mode == "answer":
                stem = re.sub(r"\s+", " ", q.text).strip()
                stem = stem if len(stem) <= 150 else stem[:147].rsplit(" ", 1)[0] + "…"
                p = b.para(f"**{n}.**\u2002{stem}", left=8, hanging=8, keep_next=True)
                b.marks_tab(p, q.marks)
                meta = f"{q.topic} · {q.difficulty} · {q.bloom}" if q.topic else f"{q.difficulty} · {q.bloom}"
                b.para(meta, italic=True, color=GREY, size=9, left=8, space_after=1)
                if q.sub_questions:
                    for s in q.sub_questions:
                        b.para(f"{s.label} [{fmt_marks(s.marks)}]", left=8, size=10)
                        b.answer(s.answer, s.options, s.marks)
                else:
                    b.answer(q.answer, q.options, q.marks)
                if q.alternative:
                    b.para(lab.get("or"), align=center, bold=True, space_before=2)
                    if q.alternative.sub_questions:
                        for s in q.alternative.sub_questions:
                            b.para(f"{s.label} [{fmt_marks(s.marks)}]", left=8, size=10)
                            b.answer(s.answer, s.options, s.marks)
                    else:
                        b.answer(q.alternative.answer, q.alternative.options, q.marks)
            else:
                show = mode == "combined"
                b.question_body(q, n, q.marks, show)
                if q.alternative:
                    b.para(lab.get("or"), align=center, bold=True, space_before=2, keep_next=True)
                    b.question_body(q.alternative, None, None, show)
            b.para(space_after=3)

    label = lab.get({"question": "question_paper", "answer": "answer_key_short", "combined": "teachers_copy_short"}[mode])
    b.footer(f"{m.subject} · {label}", f"{lab.get('set_short')} {m.set_label}" if m.set_label else "ExamForge")
    indic.complex_script_pass(b.doc, languages.resolve(m.language))
    out = io.BytesIO()
    b.doc.save(out)
    return out.getvalue()
