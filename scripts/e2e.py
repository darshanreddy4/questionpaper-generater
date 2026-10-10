"""End-to-end check of a running ExamForge server, driving it only through its HTTP API.

It starts a small fake AI service (OpenAI-compatible) so papers are written through the real provider code, billing and
caching paths without any real key. Start the server on a throwaway data folder first, for example:

    EXAMFORGE_DATA_DIR=/tmp/ef_e2e EXAMFORGE_PORT=8020 EXAMFORGE_TENANCY=multi \\
    EXAMFORGE_SUPERADMIN_EMAILS=operator@e2e.test EXAMFORGE_ALLOW_PRIVATE_AI_URLS=true \\
    EXAMFORGE_PRICE_INPUT_PER_M=3 EXAMFORGE_PRICE_OUTPUT_PER_M=15 EXAMFORGE_PLATFORM_FEE_MONTHLY=10 \\
    OPENAI_API_KEY=sk-platform-e2e OPENAI_BASE_URL=http://127.0.0.1:8099/v1 python run.py

then:  python scripts/e2e.py http://127.0.0.1:8020
"""
import io
import json
import re
import sys
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import fitz
import httpx
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tests"))
from conftest import BIOLOGY_TEXT  # noqa: E402

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8020"
FAKE_PORT = 8099
PW = "e2e-correct-horse-battery"
OPERATOR = "operator@e2e.test"
PLATFORM_KEY, BYOK_KEY = "sk-platform-e2e", "sk-byok-e2e-111111"

# ----------------------------------------------------------------------------------------------- fake AI service
USAGE = {}  # api key -> [prompt tokens, completion tokens, calls]
HINDI = ["पत्ती", "जड़", "तना", "फूल", "बीज", "फल", "मिट्टी", "पानी", "हवा", "धूप", "कोशिका", "ऊतक", "अंग", "रक्त", "हृदय", "फेफड़े"]
KANNADA = ["ಎಲೆ", "ಬೇರು", "ಕಾಂಡ", "ಹೂವು", "ಬೀಜ", "ಹಣ್ಣು", "ಮಣ್ಣು", "ನೀರು", "ಗಾಳಿ", "ಬಿಸಿಲು", "ಕೋಶ", "ಅಂಗಾಂಶ", "ರಕ್ತ", "ಹೃದಯ"]
HINDI = HINDI + [w + "ों" for w in HINDI] + [w + "ें" for w in HINDI]  # enough distinct words that no two questions look alike
KANNADA = KANNADA + [w + "ಗಳು" for w in KANNADA] + [w + "ಗಳ" for w in KANNADA]
WORDS = ["xylem", "phloem", "stomata", "chloroplast", "mitochondria", "enzyme", "membrane", "nucleus", "ribosome", "vacuole",
         "osmosis", "diffusion", "pigment", "glucose", "starch", "lactate"]


def _question(slot, system, excerpts):
    n, typ, marks = slot["slot"], slot["type"], slot["marks"]
    if "Kannada" in system:
        a, b, c = (KANNADA[(n * 3 + k) % len(KANNADA)] for k in range(3))
        text, ans = f"{a} ಮತ್ತು {b} ನಡುವಿನ {c} ಸಂಬಂಧವನ್ನು ವಿವರಿಸಿ ({n + 1}).", "ಉತ್ತರ"
        opts = [f"{KANNADA[(n * 4 + k) % len(KANNADA)]} ಆಯ್ಕೆ {k + 1}" for k in range(4)]
    elif "Hindi" in system:
        a, b, c = (HINDI[(n * 3 + k) % len(HINDI)] for k in range(3))
        text, ans = f"{a} और {b} के बीच {c} का संबंध समझाइए ({n + 1})।", "उत्तर"
        opts = [f"{HINDI[(n * 4 + k) % len(HINDI)]} विकल्प {k + 1}" for k in range(4)]
    else:
        sents = [s.strip() for s in re.split(r"(?<=[.!?])\s+", excerpts) if 40 < len(s.strip()) < 220]
        if sents:
            text = f"With reference to the passage, explain: {sents[(n * 3) % len(sents)]} (Q{n + 1})"
        else:
            a, b = WORDS[(n * 2) % len(WORDS)], WORDS[(n * 2 + 1) % len(WORDS)]
            text = f"Describe how {a} and {b} differ in structure and function (Q{n + 1})."
        ans = "Model answer"
        opts = [f"{WORDS[(n * 4 + k) % len(WORDS)]} option {k + 1}" for k in range(4)]
    item = {"slot": n, "text": text, "answer": {"text": f"{ans}: {text}"}}
    if typ == "mcq":
        item["options"], item["answer"] = opts, {"correct_option": 1, "text": ans}
    else:
        item["answer"]["steps"] = [{"text": f"step {i + 1} ({n})", "marks": marks / 2} for i in range(2)]
    return item


class Fake(BaseHTTPRequestHandler):
    def log_message(self, *a):  # quiet
        pass

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        key = self.headers.get("Authorization", "").replace("Bearer ", "")
        system = next((m["content"] for m in body["messages"] if m["role"] == "system"), "")
        user = next((m["content"] for m in body["messages"] if m["role"] == "user"), "")
        if isinstance(user, list):
            user = " ".join(p.get("text", "") for p in user)
        if "translate" in system.lower():
            out = {"texts": ["अनुवादित: " + t if "Hindi" in system else "ಅನುವಾದ: " + t for t in json.loads(user)]}
        elif "SLOTS TO WRITE" in user:
            slots = json.loads(re.search(r"SLOTS TO WRITE:\n(\[.*\])", user).group(1))
            excerpts = user.split("SOURCE EXCERPTS", 1)[-1]
            out = {"questions": [_question(s, system, excerpts) for s in slots]}
        else:
            out = {"is_diagram": False, "text": "", "figures": []}
        content = json.dumps(out, ensure_ascii=False)
        pt, ct = max(1, (len(system) + len(user)) // 4), max(1, len(content) // 3)
        u = USAGE.setdefault(key, [0, 0, 0])
        u[0] += pt
        u[1] += ct
        u[2] += 1
        payload = json.dumps({"choices": [{"message": {"role": "assistant", "content": content}}],
                              "usage": {"prompt_tokens": pt, "completion_tokens": ct}}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)


def usage(key):
    p, c, n = USAGE.get(key, [0, 0, 0])
    return p + c


# ----------------------------------------------------------------------------------------------- tiny test harness
RESULTS = []
http = httpx.Client(base_url=BASE, timeout=120)


def check(name, cond, detail=""):
    RESULTS.append((bool(cond), name))
    print(("  PASS  " if cond else "  FAIL  ") + name + (f"   [{detail}]" if not cond and detail else ""))
    return bool(cond)


def section(title):
    print(f"\n== {title}")


def H(token):
    return {"Authorization": "Bearer " + token}


def register(email, name, **kw):
    r = http.post("/api/auth/register", json={"email": email, "password": PW, "name": name, **kw})
    assert r.status_code == 200, (email, r.text)
    return H(r.json()["token"]), r.json()["user"]


def login(email):
    r = http.post("/api/auth/login", json={"email": email, "password": PW})
    return r


def upload(h, name, data, mime="text/plain"):
    r = http.post("/api/documents", headers=h, files=[("files", (name, data, mime))])
    assert r.status_code == 202, r.text
    doc = r.json()["documents"][0]
    end = time.time() + 60
    while time.time() < end:
        d = http.get(f"/api/documents/{doc['id']}", headers=h).json()
        if d["status"] in ("ready", "failed"):
            return d
        time.sleep(0.3)
    raise AssertionError("document never finished: " + name)


def make_paper(h, **body):
    body = {"subject": "Biology", "pattern_id": "unit-test-20", "diagram_mode": "none", **body}
    r = http.post("/api/papers", headers=h, json=body)
    assert r.status_code == 202, r.text
    end = time.time() + 120
    while time.time() < end:
        p = http.get(f"/api/papers/{r.json()['id']}", headers=h).json()
        if p["status"] in ("ready", "failed"):
            return p
        time.sleep(0.4)
    raise AssertionError("paper never finished")


def pdf_ok(data):
    doc = fitz.open(stream=data, filetype="pdf")
    return doc.page_count >= 1, " ".join(f[3] for p in doc for f in p.get_fonts())


def has_indic(s):
    return any("\u0900" <= ch <= "\u0d7f" for ch in s)


# ----------------------------------------------------------------------------------------------- the run
def main():
    server = ThreadingHTTPServer(("127.0.0.1", FAKE_PORT), Fake)
    threading.Thread(target=server.serve_forever, daemon=True).start()

    section("Platform operator and schools")
    op, op_user = register(OPERATOR, "Platform Operator")
    check("operator account is a platform operator", op_user["is_superadmin"])
    made = http.post("/api/platform/orgs", headers=op, json={"name": "Greenfield High", "admin_email": "head@green.test"}).json()
    g_admin, g_user = register("head@green.test", "Gita Head", invite_code=made["admin_invite_code"])
    check("operator onboards a school; its first administrator joins with the invitation", g_user["role"] == "admin" and g_user["organization"] == "Greenfield High")
    r_admin, r_user = register("head@river.test", "Ravi Head", school_name="Riverside College")
    check("a second school registers itself", r_user["role"] == "admin" and r_user["org_id"] != g_user["org_id"])
    teachers = {}
    for key, email, name, admin in (("t1", "anita@green.test", "Anita Teacher", g_admin), ("t2", "arun@green.test", "Arun Teacher", g_admin),
                                    ("r1", "bina@river.test", "Bina Teacher", r_admin)):
        code = http.post("/api/admin/invites", headers=admin, json={"email": email}).json()["code"]
        teachers[key] = register(email, name, invite_code=code)[0]
    t1, t2, r1 = teachers["t1"], teachers["t2"], teachers["r1"]
    check("teachers join their school by invitation", True)
    check("a teacher cannot create invitations or open administration", http.post("/api/admin/invites", headers=t1, json={}).status_code == 403
          and http.get("/api/admin/users", headers=t1).status_code == 403)
    check("school administrators cannot open the Platform console", http.get("/api/platform/orgs", headers=g_admin).status_code == 403
          and http.post("/api/platform/orgs", headers=g_admin, json={"name": "Rival"}).status_code == 403)
    check("a bad invitation or reused code is refused", http.post("/api/auth/register", json={"email": "x@y.test", "password": PW, "name": "X", "invite_code": made["admin_invite_code"]}).status_code == 422)
    index = http.get("/")
    check("security headers are sent", "default-src 'self'" in index.headers.get("content-security-policy", "") and index.headers.get("x-content-type-options") == "nosniff")

    section("Source files: upload, formats, languages")
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    buf = io.BytesIO()
    SimpleDocTemplate(buf).build([Paragraph("Chapter 1: Osmosis", getSampleStyleSheet()["Heading1"]),
                                  Paragraph("Osmosis is the movement of water molecules through a semipermeable membrane from a region of higher water concentration. " * 12,
                                            getSampleStyleSheet()["BodyText"])])
    d = Document()
    d.add_heading("Chapter 1: Digestion", 1)
    d.add_paragraph("Digestion is the process by which food is broken down into simple soluble substances that the body can absorb into the blood. " * 12)
    dbuf = io.BytesIO()
    d.save(dbuf)
    kannada = ("# ಅಧ್ಯಾಯ 1: ದ್ಯುತಿಸಂಶ್ಲೇಷಣೆ\n\n" + "ದ್ಯುತಿಸಂಶ್ಲೇಷಣೆ ಎಂದರೆ ಹಸಿರು ಸಸ್ಯಗಳು ಸೂರ್ಯನ ಬೆಳಕು, ನೀರು ಮತ್ತು ಕಾರ್ಬನ್ ಡೈಆಕ್ಸೈಡ್ ಬಳಸಿ ಆಹಾರ ತಯಾರಿಸುವ ಪ್ರಕ್ರಿಯೆ. ಹರಿತ್ತು ಎಲೆಗಳಲ್ಲಿರುವ ಹಸಿರು ವರ್ಣದ್ರವ್ಯ.\n" * 8).encode()
    hindi = ("# अध्याय 1: प्रकाश संश्लेषण\n\n" + "प्रकाश संश्लेषण वह प्रक्रिया है जिसमें हरे पौधे सूर्य के प्रकाश, जल और कार्बन डाइऑक्साइड से अपना भोजन बनाते हैं। क्लोरोफिल पत्तियों में पाया जाने वाला हरा वर्णक है।\n" * 8).encode()
    docs = {
        "english_txt": upload(t1, "plants.txt", BIOLOGY_TEXT.encode()),
        "pdf": upload(t1, "osmosis.pdf", buf.getvalue(), "application/pdf"),
        "docx": upload(g_admin, "digestion.docx", dbuf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "kannada": upload(g_admin, "kannada_science.txt", kannada),
        "hindi": upload(t2, "hindi_science.txt", hindi),
    }
    check("TXT, PDF, DOCX, Kannada and Hindi uploads all become ready", all(x["status"] == "ready" for x in docs.values()), [x["message"] for x in docs.values() if x["status"] != "ready"])
    check("chapters are found in English and in Kannada", len(docs["english_txt"]["chapters"]) >= 3 and docs["kannada"]["chapters"][0]["title"].startswith("ಅಧ್ಯಾಯ"))
    check("the script of each document is recognised", docs["kannada"]["stats"].get("language_hint") == "Kannada" and docs["hindi"]["stats"].get("language_hint") == "Hindi")

    section("Sharing inside a school, privacy between schools")
    ids = {x["id"] for x in docs.values()}
    for who, h in (("administrator", g_admin), ("teacher 1", t1), ("teacher 2", t2)):
        got = {x["id"] for x in http.get("/api/documents", headers=h).json()["documents"]}
        check(f"{who} of Greenfield sees all five source files", got == ids)
    check("Riverside sees none of Greenfield's files", http.get("/api/documents", headers=r1).json()["documents"] == []
          and all(http.get(f"/api/documents/{i}", headers=r_admin).status_code == 404 for i in ids))
    check("teachers cannot delete source files; another school's administrator gets 'not found'",
          http.delete(f"/api/documents/{docs['english_txt']['id']}", headers=t1).status_code == 403
          and http.delete(f"/api/documents/{docs['pdf']['id']}", headers=t2).status_code == 403
          and http.delete(f"/api/documents/{docs['pdf']['id']}", headers=r_admin).status_code == 404)

    section("Writing papers through the AI service (platform key)")
    before = usage(PLATFORM_KEY)
    p1 = make_paper(t1, document_ids=[docs["english_txt"]["id"]], seed=11, class_level="9", board="State", institution="Greenfield High", exam_name="Unit Test 1")
    check("teacher 1 generates a Biology paper from the English notes", p1["status"] == "ready", p1["message"])
    q = p1["quality"]
    check("every question has an answer and none failed", q["failed"] == [] and q["missing_answers"] == [], q["failed"])
    qs = [x for s in p1["content"]["sections"] for x in s["questions"]]
    check("questions are built from the uploaded passages", sum("With reference to the passage" in x["text"] for x in qs) >= 4)
    check("the paper used AI tokens that the provider reported", usage(PLATFORM_KEY) > before and q["usage"]["prompt_tokens"] > 0)

    p2 = make_paper(t2, subject="Mathematics", seed=12, syllabus_text="Quadratic equations")
    check("teacher 2 generates a Mathematics paper (code writes what it can, AI the rest)", p2["status"] == "ready" and p2["quality"]["failed"] == [], p2["message"])
    check("maths answers were checked by computer, not just trusted", any(x.get("verified") in ("checked", "construction") for s in p2["content"]["sections"] for x in s["questions"]))

    p3 = make_paper(g_admin, subject="विज्ञान", language="Hindi", seed=13, document_ids=[docs["hindi"]["id"]], institution="ग्रीनफ़ील्ड हाई स्कूल")
    check("the administrator generates a Hindi paper from the Hindi file", p3["status"] == "ready" and p3["quality"]["failed"] == [], p3["message"])
    hq = [x for s in p3["content"]["sections"] for x in s["questions"]]
    check("every Hindi question is in Devanagari", all(has_indic(x["text"]) for x in hq))
    check("the pattern's own wording was translated", p3["content"]["meta"]["instructions"][0].startswith("अनुवादित"))

    p4 = make_paper(g_admin, subject="ಕನ್ನಡ", language="Kannada", seed=14, document_ids=[docs["kannada"]["id"]])
    check("a Kannada-subject paper in Kannada is generated from Kannada notes", p4["status"] == "ready" and all(has_indic(x["text"]) for s in p4["content"]["sections"] for x in s["questions"]), p4["message"])

    section("Downloads")
    for label, p, font in (("English", p1, ""), ("Maths", p2, ""), ("Hindi", p3, "NotoSansDevanagari"), ("Kannada", p4, "NotoSansKannada")):
        for mode in ("question", "answer", "combined"):
            r = http.get(f"/api/papers/{p['id']}/export", headers=g_admin, params={"format": "pdf", "mode": mode})
            ok, fonts = pdf_ok(r.content) if r.status_code == 200 else (False, "")
            check(f"{label} {mode} PDF opens and has pages" + (f" with {font}" if font else ""), r.status_code == 200 and ok and (not font or font in fonts), r.text[:100])
        r = http.get(f"/api/papers/{p['id']}/export", headers=g_admin, params={"format": "docx", "mode": "combined"})
        words = "\n".join(x.text for x in Document(io.BytesIO(r.content)).paragraphs) if r.status_code == 200 else ""
        check(f"{label} DOCX opens and has the paper text", r.status_code == 200 and len(words) > 200, r.text[:100])
    docx_h = "\n".join(x.text for x in Document(io.BytesIO(http.get(f"/api/papers/{p3['id']}/export", headers=g_admin, params={"format": "docx"}).content)).paragraphs)
    check("Hindi paper headings print in Hindi", "समय" in docx_h and "Time Allowed" not in docx_h)
    z = http.get(f"/api/papers/{p1['id']}/export-all", headers=t2, params={"sets": 2})
    names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    check("two answer-set variants download as a zip of PDF and DOCX files", len(names) == 8 and any("setB" in n for n in names), names)

    section("Editing rules")
    content = p1["content"]
    content["sections"][0]["questions"][0]["text"] = "Edited by the creator: define osmosis and give one example from plants?"
    check("the creator can edit", http.put(f"/api/papers/{p1['id']}", headers=t1, json=content).status_code == 200)
    check("a colleague can read but not edit, delete or regenerate",
          http.get(f"/api/papers/{p1['id']}", headers=t2).json()["can_edit"] is False
          and http.put(f"/api/papers/{p1['id']}", headers=t2, json=content).status_code == 403
          and http.delete(f"/api/papers/{p1['id']}", headers=t2).status_code == 403
          and http.post(f"/api/papers/{p1['id']}/questions/q1/regenerate", headers=t2, json={}).status_code == 403)
    rg = http.post(f"/api/papers/{p1['id']}/questions/q6/regenerate", headers=t1, json={"instruction": "make it harder"})
    check("the creator regenerates one question through the AI", rg.status_code == 200 and rg.json().get("question", {}).get("text"), rg.text[:150])
    check("the administrator can edit a teacher's paper", http.put(f"/api/papers/{p1['id']}", headers=g_admin, json=content).status_code == 200)

    section("Question bank")
    bank = http.get("/api/bank", headers=t2).json()
    check("everyone in the school sees the bank, with who wrote each question", bank["total"] > 10 and {i["created_by"] for i in bank["items"]} >= {"Anita Teacher"})
    check("Riverside's bank is empty", http.get("/api/bank", headers=r1).json()["total"] == 0)
    plan = http.post("/api/papers/blueprint", headers=t2, json={"subject": "Biology", "pattern_id": "unit-test-20", "seed": 11, "reuse_bank": True, "document_ids": [docs["english_txt"]["id"]]}).json()
    check("a colleague's new paper can reuse the bank; the cost preview says so", plan["writers"]["bank"] > 0, plan["writers"])
    item = bank["items"][0]["id"]
    check("only the administrator can delete a bank question", http.delete(f"/api/bank/{item}", headers=t1).status_code == 403
          and http.delete(f"/api/bank/{item}", headers=r_admin).status_code == 404 and http.delete(f"/api/bank/{item}", headers=g_admin).status_code == 204)

    section("Billing: platform key, own key, per-school rates")
    bill = http.get("/api/admin/billing", headers=g_admin).json()["summary"]
    greenfield_platform = usage(PLATFORM_KEY)
    check("platform tokens billed to the school equal what the AI service reported", bill["platform"]["tokens"] == greenfield_platform, (bill["platform"]["tokens"], greenfield_platform))
    check("fee plus token charge make the invoice total", abs(bill["total"] - (bill["platform_fee"] + bill["platform"]["charge"])) < 1e-6 and bill["platform_fee"] == 10.0)
    check("a teacher sees only their own usage", http.get("/api/billing", headers=t1).json()["platform"]["tokens"] < bill["platform"]["tokens"])
    ok_key = http.put("/api/ai/settings", headers=t1, json={"scope": "user", "provider": "openai_compatible", "model": "m1", "base_url": f"http://127.0.0.1:{FAKE_PORT}/v1", "api_key": BYOK_KEY})
    check("a teacher saves their own AI key (stored encrypted, shown masked)", ok_key.status_code == 200 and BYOK_KEY not in http.get("/api/ai/settings", headers=t1).text)
    check("the key's connection test passes", http.post("/api/ai/test", headers=t1, json={"scope": "user"}).json().get("ok") is True)
    plat_before, byok_before = usage(PLATFORM_KEY), usage(BYOK_KEY)
    p5 = make_paper(t1, seed=15, document_ids=[docs["pdf"]["id"]])
    check("a paper on the teacher's own key goes to the teacher's AI account, not the platform's", p5["status"] == "ready" and usage(BYOK_KEY) > byok_before and usage(PLATFORM_KEY) == plat_before)
    mine = http.get("/api/billing", headers=t1).json()
    check("own-key tokens are shown but never charged", mine["own_key"]["tokens"] == usage(BYOK_KEY) - byok_before and mine["own_key"]["charge"] == 0, (mine["own_key"], usage(BYOK_KEY) - byok_before))
    bill2 = http.get("/api/admin/billing", headers=g_admin).json()["summary"]
    check("the school's platform charge did not move for the own-key paper", bill2["platform"]["tokens"] == bill["platform"]["tokens"])
    http.delete("/api/ai/settings", headers=t1, params={"scope": "user"})
    orgs = {o["name"]: o for o in http.get("/api/platform/orgs", headers=op).json()["orgs"]}
    check("the operator's Platform page shows the same tokens and total for the school",
          orgs["Greenfield High"]["month"]["platform_tokens"] == bill["platform"]["tokens"] and abs(orgs["Greenfield High"]["month"]["total"] - bill["total"]) < 1e-6)
    http.patch(f"/api/platform/orgs/{g_user['org_id']}", headers=op, json={"platform_fee_monthly": 50.0, "price_input_per_million": 100.0})
    p6 = make_paper(t2, subject="Botany", seed=16, syllabus_text="Roots")
    bill3 = http.get("/api/admin/billing", headers=g_admin).json()["summary"]
    check("a negotiated fee and token price apply to that school only", bill3["platform_fee"] == 50.0 and bill3["pricing"]["input_per_million"] == 100.0
          and http.get("/api/admin/billing", headers=r_admin).json()["summary"]["platform_fee"] == 10.0)

    section("Platform support access")
    check("support cannot open a school that has not allowed it", http.post(f"/api/platform/orgs/{g_user['org_id']}/support-session", headers=op).status_code == 403)
    check("only the school administrator can allow it", http.post("/api/admin/support", headers=t1, json={"hours": 1}).status_code == 403
          and http.post("/api/admin/support", headers=g_admin, json={"hours": 2, "note": "help with Section B"}).status_code == 201)
    sess = http.post(f"/api/platform/orgs/{g_user['org_id']}/support-session", headers=op)
    sh = H(sess.json()["token"])
    check("with permission the operator can look", sess.status_code == 201 and len(http.get("/api/papers", headers=sh).json()["papers"]) >= 5)
    check("the support view is read-only and content-only",
          http.delete(f"/api/papers/{p3['id']}", headers=sh).status_code == 403 and http.get("/api/admin/users", headers=sh).status_code == 403
          and http.get(f"/api/papers/{p3['id']}/export", headers=sh).status_code == 403 and http.get("/api/ai/settings", headers=sh).status_code == 403
          and http.get(f"/api/papers/{p3['id']}", headers=sh).status_code == 200)
    check("a support session on Greenfield cannot see Riverside", http.get("/api/papers", headers=sh).json()["papers"] != [] and
          not any(x["created_by"] == "Bina Teacher" for x in http.get("/api/papers", headers=sh).json()["papers"]))
    http.delete("/api/admin/support", headers=g_admin)
    check("stopping support cuts off the open session at once", http.get("/api/papers", headers=sh).status_code == 403)
    actions = [e["action"] for e in http.get("/api/admin/audit?limit=300", headers=g_admin).json()["events"]]
    check("the school's audit log records the permission, the session and what was opened", all(a in actions for a in ("support.allowed", "support.session_start", "support.view", "support.revoked")))
    check("Riverside's log shows none of it", not any(e["action"].startswith("support.") for e in http.get("/api/admin/audit", headers=r_admin).json()["events"]))

    section("Isolation and account controls")
    gid = p3["id"]
    check("Riverside cannot read, change, delete or download Greenfield's papers",
          all(http.request(m, f"/api/papers/{gid}{suffix}", headers=r_admin, **kw).status_code == 404
              for m, suffix, kw in (("GET", "", {}), ("DELETE", "", {}), ("GET", "/export", {}), ("PUT", "", {"json": p3["content"]}))))
    check("each school's audit log only holds its own people",
          not ({e["user_id"] for e in http.get("/api/admin/audit?limit=300", headers=g_admin).json()["events"]} & {e["user_id"] for e in http.get("/api/admin/audit", headers=r_admin).json()["events"]}))
    http.patch(f"/api/platform/orgs/{r_user['org_id']}", headers=op, json={"is_active": False})
    check("a suspended school is locked out; others are not", http.get("/api/papers", headers=r_admin).status_code == 403
          and login("bina@river.test").status_code == 403 and http.get("/api/papers", headers=g_admin).status_code == 200)
    http.patch(f"/api/platform/orgs/{r_user['org_id']}", headers=op, json={"is_active": True})
    check("reactivating restores access", http.get("/api/papers", headers=r_admin).status_code == 200)
    tid = next(u["id"] for u in http.get("/api/admin/users", headers=g_admin).json()["users"] if u["email"] == "arun@green.test")
    http.patch(f"/api/admin/users/{tid}", headers=g_admin, json={"is_active": False})
    check("a disabled teacher cannot sign in or use an old session", login("arun@green.test").status_code in (401, 403) and http.get("/api/papers", headers=t2).status_code in (401, 403))
    check("deleting a source file as administrator removes it for everyone", http.delete(f"/api/documents/{docs['pdf']['id']}", headers=g_admin).status_code == 204
          and http.get(f"/api/documents/{docs['pdf']['id']}", headers=t1).status_code == 404)
    check("the bad-password limiter answers 429 after repeated attempts", any(http.post("/api/auth/login", json={"email": "nobody@x.test", "password": "wrong-password-1"}).status_code == 429 for _ in range(15)))

    passed = sum(1 for ok, _ in RESULTS if ok)
    print(f"\n{passed} of {len(RESULTS)} checks passed")
    for ok, name in RESULTS:
        if not ok:
            print("  FAILED:", name)
    server.shutdown()
    sys.exit(0 if passed == len(RESULTS) else 1)


if __name__ == "__main__":
    main()
