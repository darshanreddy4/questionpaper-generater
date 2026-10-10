"""The file reader is bundled in this project (book_to_skill/), so ExamForge runs without any sibling folder."""
import sys
from pathlib import Path

import pytest

import book_to_skill
from app.config import load_settings
from app.ingestion.extract import extract_source
from app.llm import OfflineProvider

ROOT = Path(__file__).resolve().parent.parent
BODY = ("Photosynthesis is the process by which green plants prepare food using sunlight, water and carbon dioxide. "
        "Chlorophyll absorbs the light energy that drives the reaction inside the chloroplast of every leaf. ") * 6


def _read(path: Path) -> str:
    return extract_source(path, OfflineProvider(), load_settings(load_env_file=False)).text


def test_the_reader_comes_from_this_project_not_from_a_sibling_folder():
    assert Path(book_to_skill.__file__).resolve().is_relative_to(ROOT) if hasattr(Path, "is_relative_to") else str(ROOT) in book_to_skill.__file__
    assert not any("book-to-skill-master" in p for p in sys.path)
    assert (ROOT / "book_to_skill" / "LICENSE.md").read_text().startswith("MIT License")  # its licence travels with the code
    assert not any(line.strip().startswith("-e ../") for line in (ROOT / "requirements.txt").read_text().splitlines())


@pytest.mark.parametrize("ext", [".txt", ".md", ".html", ".rtf", ".docx", ".pdf", ".epub"])
def test_every_common_format_can_be_read(tmp_path, ext):
    path = tmp_path / f"notes{ext}"
    if ext in (".txt", ".md"):
        path.write_text("# Chapter 1: Photosynthesis\n\n" + BODY, encoding="utf-8")
    elif ext == ".html":
        path.write_text(f"<html><body><h1>Chapter 1: Photosynthesis</h1><p>{BODY}</p></body></html>", encoding="utf-8")
    elif ext == ".rtf":
        path.write_text("{\\rtf1\\ansi\\deff0 {\\b Chapter 1: Photosynthesis}\\par " + BODY + "}", encoding="utf-8")
    elif ext == ".docx":
        from docx import Document
        doc = Document()
        doc.add_heading("Chapter 1: Photosynthesis", 1)
        doc.add_paragraph(BODY)
        doc.save(str(path))
    elif ext == ".pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.platypus import Paragraph, SimpleDocTemplate
        from reportlab.lib.styles import getSampleStyleSheet
        SimpleDocTemplate(str(path), pagesize=A4).build([Paragraph("Chapter 1: Photosynthesis", getSampleStyleSheet()["Heading1"]),
                                                         Paragraph(BODY, getSampleStyleSheet()["BodyText"])])
    else:
        from ebooklib import epub
        book = epub.EpubBook()
        book.set_identifier("t1"); book.set_title("Notes"); book.set_language("en")
        chapter = epub.EpubHtml(title="Photosynthesis", file_name="c1.xhtml", lang="en")
        chapter.content = f"<h1>Chapter 1: Photosynthesis</h1><p>{BODY}</p>"
        book.add_item(chapter)
        book.toc = (chapter,)
        book.add_item(epub.EpubNcx()); book.add_item(epub.EpubNav())
        book.spine = ["nav", chapter]
        epub.write_epub(str(path), book)
    text = _read(path)
    assert "chlorophyll" in text.lower() and "photosynthesis" in text.lower(), (ext, text[:200])
