import dataclasses
import io
import json
import re

import pytest
from fastapi.testclient import TestClient

from conftest import register
from app.exam.generator import PaperGenerator
from app.exam.hybrid_engine import CodeOnlyEngine
from app.exam.llm_engine import LLMEngine
from app.exam.offline_engine import OfflineEngine
from app.exam.patterns import get_preset
from app.exam.schema import GenerationRequest, TopicSpec
from app.llm.base import LLMProvider
from app.llm.usage import UsageMeter
from app.main import create_app
from app.models import AiSetting, UsageLog
from app.services.ai_settings import decrypt_key, encrypt_key, resolve_ai
from app.services.usage import billing_summary, record_usage

KEY = "sk-test-1234567890abcd"


class FakeModel(LLMProvider):
    name, model = "fake", "fake-1"
    supports_vision = True

    def __init__(self):
        super().__init__(2)
        self.calls = 0
        self.slots = []

    def _complete(self, system, user, images, json_mode, max_tokens, temperature):
        self.calls += 1
        m = re.search(r"SLOTS TO WRITE:\n(\[.*\])", user)
        if m:
            self.slots += [s["slot"] for s in json.loads(m.group(1))]
        return '{"is_diagram": false}'


def _client(settings, **changes):
    app = create_app(dataclasses.replace(settings, **changes))
    return app, TestClient(app)


def _put(client, h, **kw):
    body = {"scope": "user", "provider": "openai", "model": "gpt-4o", "api_key": KEY}
    body.update(kw)
    return client.put("/api/ai/settings", headers=h, json=body)


# ----------------------------------------------------------------- keys

def test_keys_are_encrypted_masked_and_used_for_the_owner_only(settings):
    app, c = _client(settings)
    with c:
        admin = register(c)
        teacher = register(c, "t2@example.com")
        assert _put(c, teacher).status_code == 200
        view = c.get("/api/ai/settings", headers=teacher)
        assert KEY not in view.text and view.json()["user"]["key_hint"] == "abcd"
        assert view.json()["effective"]["source"] == "user" and view.json()["effective"]["billing"] == "byok"
        with app.state.ctx.db.session() as db:
            row = db.query(AiSetting).one()
            assert KEY not in row.key_cipher and decrypt_key(app.state.ctx.settings, row.key_cipher) == KEY
        assert c.get("/api/ai/settings", headers=admin).json()["effective"]["source"] == "offline"  # not shared
        r = c.delete("/api/ai/settings?scope=user", headers=teacher)
        assert r.status_code == 204
        assert c.get("/api/ai/settings", headers=teacher).json()["effective"]["source"] == "offline"


def test_key_can_be_kept_when_only_the_model_changes(settings):
    app, c = _client(settings)
    with c:
        h = register(c)
        _put(c, h)
        assert _put(c, h, model="gpt-4o-mini", api_key="").status_code == 200
        with app.state.ctx.db.session() as db:
            row = db.query(AiSetting).one()
            assert row.model == "gpt-4o-mini" and decrypt_key(app.state.ctx.settings, row.key_cipher) == KEY
        assert _put(c, register(c, "x@example.com"), api_key="").status_code == 422  # a new setting needs a key


def test_custom_endpoints_cannot_reach_internal_networks(settings):
    _, c = _client(settings)
    with c:
        h = register(c)
        for url in ("https://127.0.0.1/v1", "https://169.254.169.254/latest", "https://10.0.0.5/v1",
                    "http://example.com/v1", "https://user:pw@example.com/v1", "not a url"):
            r = _put(c, h, provider="openai_compatible", base_url=url)
            assert r.status_code == 422, url
    _, c2 = _client(settings, allow_private_ai_urls=True)
    with c2:
        h2 = register(c2, "other@example.com")
        assert _put(c2, h2, provider="openai_compatible", model="llama3", base_url="http://127.0.0.1:11434/v1").status_code == 200


def test_school_key_is_admin_only_and_teachers_inherit_it_unless_they_set_their_own(settings):
    _, c = _client(settings)
    with c:
        admin = register(c)
        teacher = register(c, "t2@example.com")
        assert _put(c, teacher, scope="workspace").status_code == 403
        assert _put(c, admin, scope="workspace", model="school-model").status_code == 200
        seen = c.get("/api/ai/settings", headers=teacher).json()
        assert seen["effective"]["source"] == "workspace" and seen["workspace"] == {"configured": True}
        assert "key_hint" not in json.dumps(seen["workspace"])  # details stay with administrators
        _put(c, teacher, model="own-model")
        assert c.get("/api/ai/settings", headers=teacher).json()["effective"]["model"] == "own-model"


def test_testing_a_saved_key_never_echoes_it(settings, monkeypatch):
    app, c = _client(settings)
    with c:
        h = register(c)
        _put(c, h)
        from app.llm import providers

        def boom(self, *a, **k):
            from app.llm.base import LLMError
            raise LLMError(f"Model API error 401: bad key {KEY}")

        monkeypatch.setattr(providers.OpenAICompatProvider, "_complete", boom)
        r = c.post("/api/ai/test", headers=h, json={"scope": "user"}).json()
        assert r["ok"] is False and KEY not in json.dumps(r)


# ----------------------------------------------------------------- billing

def test_only_platform_tokens_are_charged_and_zero_usage_means_fee_only(settings):
    app, c = _client(settings, price_input_per_million=3.0, price_output_per_million=15.0, platform_fee_monthly=10.0)
    with c:
        h = register(c)
        uid = c.get("/api/auth/me", headers=h).json()["id"]
        state = app.state.ctx
        quiet = billing_summary(state, None)
        assert quiet["total"] == 10.0 and "only the platform fee" in quiet["summary"]

        meter = UsageMeter()
        meter.add(1000, 500)
        state.provider = FakeModel()
        record_usage(state, uid, "paper", "p1", meter, resolve_ai(state, uid))  # platform key
        _put(c, h)
        record_usage(state, uid, "paper", "p2", meter, resolve_ai(state, uid))  # the teacher's own key

        summary = billing_summary(state, None)
        assert summary["platform"]["tokens"] == 1500 and summary["own_key"]["tokens"] == 1500
        assert summary["platform"]["charge"] == pytest.approx(0.0105)  # 1000*3 + 500*15 millionths
        assert summary["total"] == pytest.approx(10.0105)
        mine = c.get("/api/billing", headers=h).json()
        assert mine["platform_fee"] == 0 and mine["total"] == pytest.approx(0.0105)  # the flat fee is the school's
        assert c.get("/api/admin/billing", headers=h).json()["by_user"][0]["own_key_tokens"] == 1500
        with state.db.session() as db:
            assert sorted(r.billing for r in db.query(UsageLog)) == ["byok", "platform"]


def test_teachers_cannot_read_workspace_billing(settings):
    _, c = _client(settings)
    with c:
        register(c)
        teacher = register(c, "t2@example.com")
        assert c.get("/api/admin/billing", headers=teacher).status_code == 403


def test_own_key_is_exempt_from_the_platform_token_budget(settings):
    app, c = _client(settings, daily_token_budget=5)
    with c:
        h = register(c)
        uid = c.get("/api/auth/me", headers=h).json()["id"]
        app.state.ctx.provider = FakeModel()
        with app.state.ctx.db.session() as db:
            db.add(UsageLog(user_id=uid, kind="paper", prompt_tokens=50, billing="platform"))
        body = {"subject": "Mathematics", "pattern_id": "unit-test-20"}
        assert c.post("/api/papers", headers=h, json=body).status_code == 429
        _put(c, h)
        assert c.post("/api/papers", headers=h, json=body).status_code == 202


# ----------------------------------------------------------------- writing modes

def _maths():
    return GenerationRequest(subject="Mathematics", seed=3, topics=[TopicSpec(name=t) for t in (
        "Quadratic equations", "Probability", "Proofs in number theory")])


def test_code_only_mode_spends_nothing_and_leaves_ai_slots_for_the_teacher():
    content, quality, slots = PaperGenerator(CodeOnlyEngine(OfflineEngine(3)), _maths(), get_preset("unit-test-20"),
                                             workers=1).run()
    assert set(quality["writers"]) == {"code"}
    assert quality["failed"] and all("Code-only" in f["error"] for f in quality["failed"])
    text = " ".join(q.text for s in content.sections for q in s.questions)
    assert "could not be generated" in text


def test_ai_all_mode_sends_every_slot_to_the_model():
    model = FakeModel()
    PaperGenerator(LLMEngine(model), _maths(), get_preset("unit-test-20"), workers=1).run()
    assert len(set(model.slots)) == 11


def test_make_engine_follows_the_chosen_mode(settings):
    from app.exam.hybrid_engine import HybridEngine
    from app.services.ai_settings import ResolvedAI
    from app.services.papers import make_engine

    app, c = _client(settings)
    with c:
        state, resolved = app.state.ctx, ResolvedAI(FakeModel(), "platform", "platform")
        assert isinstance(make_engine(state, 1, "auto", resolved), HybridEngine)
        assert isinstance(make_engine(state, 1, "ai_all", resolved), LLMEngine)
        assert isinstance(make_engine(state, 1, "code_only", resolved), CodeOnlyEngine)
        assert isinstance(make_engine(state, 1, "ai_all", ResolvedAI(state.provider, "offline", "none")), OfflineEngine)


def test_blueprint_previews_who_will_write_what(client):
    h = register(client)
    r = client.post("/api/papers/blueprint", headers=h, json={
        "subject": "Mathematics", "pattern_id": "unit-test-20", "seed": 2, "diagram_mode": "none",
        "syllabus_text": "Quadratic equations\nProbability"}).json()["writers"]
    assert r["bank"] + r["code"] + r["ai"] + r["unwritten"] == r["questions"] == 11
    assert r["code"] > 0 and r["ai_available"] is False and r["uses_tokens"] is False


# ----------------------------------------------------------------- AI is opt-in for pictures

def _png_bytes():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (200, 120), "white").save(buf, "PNG")
    return buf.getvalue()


def test_pictures_and_scans_need_an_explicit_opt_in(client):
    h = register(client)
    r = client.post("/api/documents", headers=h, files=[("files", ("fig.png", _png_bytes(), "image/png"))])
    assert r.status_code == 202
    import time

    def failed_message(doc_id):
        for _ in range(100):
            d = client.get(f"/api/documents/{doc_id}", headers=h).json()
            if d["status"] in ("failed", "ready"):
                return d["message"]
            time.sleep(0.1)

    assert "ticked" in failed_message(r.json()["documents"][0]["id"])
    r2 = client.post("/api/documents", headers=h, data={"ai_vision": "true"},
                     files=[("files", ("fig2.png", _png_bytes(), "image/png"))])
    assert "AI & billing" in failed_message(r2.json()["documents"][0]["id"])  # opted in, but no AI available


def test_figures_in_documents_are_not_sent_to_the_ai_unless_requested(settings, tmp_path):
    from docx import Document

    from app.ingestion.extract import extract_source

    path = tmp_path / "book.docx"
    doc = Document()
    doc.add_paragraph("Photosynthesis is how plants make food. " * 20)
    img = tmp_path / "f.png"
    img.write_bytes(_png_bytes())
    doc.add_picture(str(img))
    doc.save(path)
    spy = FakeModel()
    off = extract_source(path, spy, settings, vision=False)
    assert spy.calls == 0 and off.stats["figures"] == 0
    extract_source(path, spy, settings, vision=True)
    assert spy.calls == 1
