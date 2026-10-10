import dataclasses
import io
import json
import re
import time
import zipfile

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import Paragraph as RLParagraph

from app import languages, textutil
from app.exam.generator import PaperGenerator, language_problem, resolve_topics
from app.exam.hybrid_engine import HybridEngine
from app.exam.llm_engine import LLMEngine, language_rule, token_factor
from app.exam.offline_engine import OfflineEngine
from app.exam.patterns import get_preset
from app.exam.schema import (
    Alternative, Answer, GenerationRequest, Match, PaperContent, PaperMeta, PaperSection, Question, Step,
)
from app.export import build_docx, build_pdf, indic
from app.export.labels import for_language
from app.export.mathtext import render_math
from app.ingestion.chunking import chunk_text
from app.ingestion.extract import _sanitize
from app.ingestion.retrieval import Retriever, compress_passage, tokenize
from app.llm.base import LLMProvider
from app.llm.usage import estimate_tokens
from app.main import create_app
from app.models import BankQuestion
from app.services.papers import _bank_items

from conftest import register

HINDI_PROSE = "प्रकाश संश्लेषण की प्रक्रिया में पौधे सूर्य के प्रकाश से भोजन बनाते हैं।"
TELUGU_PROSE = "కిరణజన్య సంయోగక్రియలో మొక్కలు సూర్యరశ్మి నుండి ఆహారాన్ని తయారు చేస్తాయి."

SAMPLES = {
    "devanagari": "क्लोरोफिल की भूमिका स्पष्ट कीजिए।", "bengali": "সালোকসংশ্লেষে ক্লোরোফিলের ভূমিকা ব্যাখ্যা করো।",
    "gurmukhi": "ਪ੍ਰਕਾਸ਼ ਸੰਸ਼ਲੇਸ਼ਣ ਵਿੱਚ ਕਲੋਰੋਫਿਲ ਦੀ ਭੂਮਿਕਾ ਦੱਸੋ।", "gujarati": "પ્રકાશસંશ્લેષણમાં હરિતદ્રવ્યની ભૂમિકા સમજાવો.",
    "oriya": "ସାଲୋକସଂଶ୍ଳେଷଣରେ କ୍ଲୋରୋଫିଲର ଭୂମିକା ବ୍ୟାଖ୍ୟା କର।", "tamil": "ஒளிச்சேர்க்கையில் பச்சையத்தின் பங்கை விளக்குக.",
    "telugu": TELUGU_PROSE, "kannada": "ದ್ಯುತಿಸಂಶ್ಲೇಷಣೆಯಲ್ಲಿ ಹರಿತ್ತಿನ ಪಾತ್ರವನ್ನು ವಿವರಿಸಿ.",
    "malayalam": "പ്രകാശസംശ്ലേഷണത്തിൽ ഹരിതകത്തിന്റെ പങ്ക് വിശദീകരിക്കുക.",
}
LANG_OF_SCRIPT = {"devanagari": "Hindi", "bengali": "Bengali", "gurmukhi": "Punjabi", "gujarati": "Gujarati",
                  "oriya": "Odia", "tamil": "Tamil", "telugu": "Telugu", "kannada": "Kannada", "malayalam": "Malayalam"}


def _paper(language="Hindi", text=HINDI_PROSE, **meta):
    hindi = languages.has_indic(text)
    opts = ["ऑक्सीजन", "कार्बन", "नाइट्रोजन", "हाइड्रोजन"] if hindi else ["Oxygen", "Carbon", "Nitrogen", "Hydrogen"]
    q1 = Question(id="q1", type="mcq", text=text, marks=1, options=opts, answer=Answer(text=opts[1], correct_option=1))
    steps = [Step(text=text, marks=1.5), Step(text=text, marks=1.5)]
    q2 = Question(id="q2", type="short", text=text + " $x^2$", marks=3, answer=Answer(text=text, steps=steps),
                  alternative=Alternative(text=text, answer=Answer(text=text, steps=steps)))
    q3 = Question(id="q3", type="match_following", text=text, marks=2,
                  match=Match(left=["a", "b", "c"], right=["x", "y", "z"]), answer=Answer(text="1-a"))
    m = PaperMeta(institution=text, exam_name=text, subject=text, class_level="10", language=language,
                  duration_minutes=150, total_marks=6, instructions=[text], set_label="A", **meta)
    return PaperContent(meta=m, sections=[
        PaperSection(id="A", title=text, question_type="mcq", marks_each=1, count=1, questions=[q1]),
        PaperSection(id="B", title=text, question_type="short", marks_each=3, count=1, attempt=1, negative_marks=0.25,
                     questions=[q2, q3])])


# ----------------------------------------------------------------------------- languages

def test_languages_resolve_by_name_native_code_and_phrase():
    assert languages.resolve("Hindi").code == "hi" and languages.resolve("हिन्दी").code == "hi"
    assert languages.resolve("te").name == "Telugu" and languages.resolve("Hindi medium").code == "hi"
    assert languages.resolve("Oriya").code == "or"
    assert languages.resolve("French") is None and languages.resolve("").code == "en"
    assert languages.is_english("English") and languages.is_english("English (India)") and languages.is_english(None)
    assert not languages.is_english("Hindi") and not languages.is_english("French")
    assert languages.key("Hindi") == "hi" and languages.key("French") == "french" and languages.key("") == "en"


def test_scripts_are_told_apart_and_joiners_and_danda_stay_with_the_word():
    for script, text in SAMPLES.items():
        assert languages.dominant_script(text * 3) == script, script
    assert languages.dominant_script("Plain English text only " * 10) is None
    runs = languages.runs("क्षेत्रफल। ന്\u200d x")
    assert [s for s, _ in runs if s] == ["devanagari", "malayalam"]
    assert runs[0] == ("devanagari", "क्षेत्रफल।")  # the danda joins the word before it
    assert languages.indic_share("hello there world") == 0.0 and languages.indic_share("hi") is None


# ----------------------------------------------------------------------------- reading Indic text

def test_words_keep_their_vowel_signs_and_sentences_split_at_the_danda():
    words = textutil.WORD.findall(HINDI_PROSE)
    assert words[:3] == ["प्रकाश", "संश्लेषण", "की"] and len(words) == 13
    assert textutil.WORD.findall(TELUGU_PROSE)[:2] == ["కిరణజన్య", "సంయోగక్రియలో"]
    parts = textutil.SENTENCE_SPLIT.split("पहला वाक्य यहाँ है। दूसरा वाक्य यहाँ है। Third one.")
    assert len(parts) == 3 and parts[0].endswith("।")


def test_search_and_chunking_work_on_hindi_and_telugu():
    text = ("# अध्याय 1: प्रकाश संश्लेषण\n\n" + (HINDI_PROSE + " ") * 15 + "\n\n"
            "# పాఠం 2: శ్వాసక్రియ\n\n" + (TELUGU_PROSE + " ") * 15)
    chunks = chunk_text(text)
    assert [c.chapter for c in chunks] == ["अध्याय 1: प्रकाश संश्लेषण", "పాఠం 2: శ్వాసక్రియ"]
    assert tokenize("సంయోగక్రియలో మొక్కలు")[:1] == ["సంయోగక"]  # Indic words are kept to their root

    class C:
        def __init__(self, i, chapter, text):
            self.id, self.chapter, self.text, self.kind = i, chapter, text, "text"

    r = Retriever([C(str(i), c.chapter, c.text) for i, c in enumerate(chunks)])
    assert r.search("సంయోగక్రియ మొక్కలు")[0][1].chapter.startswith("పాఠం")
    assert r.search("संश्लेषण प्रकाश पौधे")[0][1].chapter.startswith("अध्याय")
    kept = compress_passage((HINDI_PROSE + " ") * 3 + "बिल्कुल अलग विषय का वाक्य यहाँ है। " * 4, "प्रकाश संश्लेषण", 90)
    assert "प्रकाश" in kept and len(kept) <= 130


def test_zero_width_joiners_after_indic_letters_survive_sanitising_but_hidden_characters_do_not():
    text = "ന്\u200d क्\u200dष నా\u200c x\u200by"
    cleaned, removed = _sanitize(text)
    assert cleaned == "ന്\u200d क्\u200dष నా\u200c xy" and removed >= 1


def test_token_estimates_charge_indic_text_at_its_real_cost():
    assert estimate_tokens("ప్రకాశ సంశ్లేషణ " * 10) > 2.5 * estimate_tokens("hello world " * 10)
    assert token_factor("English") == 1.0 and token_factor("Telugu") > 3 and token_factor("French") == 1.5


# ----------------------------------------------------------------------------- PDF

@pytest.mark.parametrize("script", sorted(SAMPLES))
def test_pdf_embeds_the_script_font_for_every_supported_script(script):
    lang = LANG_OF_SCRIPT[script]
    data = build_pdf(_paper(lang, SAMPLES[script]), "combined")
    page_fonts = " ".join(f[3] for p in fitz.open(stream=data, filetype="pdf") for f in p.get_fonts())
    assert "Noto" in page_fonts, (script, page_fonts)
    assert data.startswith(b"%PDF")


def test_english_pdf_is_unchanged_and_embeds_no_indic_font():
    data = build_pdf(_paper("English", "Define photosynthesis."), "question")
    fonts = " ".join(f[3] for p in fitz.open(stream=data, filetype="pdf") for f in p.get_fonts())
    assert "Noto" not in fonts
    text = "\n".join(p.get_text() for p in fitz.open(stream=data, filetype="pdf"))
    assert "Time Allowed" in text and "SECTION A" in text and "Maximum Marks" in text


def test_headings_print_in_the_papers_language_and_numbers_stay_ascii():
    text = "\n".join(p.get_text() for p in fitz.open(stream=build_pdf(_paper("Hindi"), "question"), filetype="pdf"))
    assert "Time Allowed" not in text and "SECTION" not in text and "Page " not in text
    assert "2 घंटे 30 मिनट".replace("घंटे", "घंटे") and "150" not in text
    lab = for_language("Hindi")
    assert lab.duration(150) == "2 घंटे 30 मिनट" and lab.get("page", p=1, t=3) == "पृष्ठ 1 / 3"
    assert for_language("Nepali").get("time_allowed") == "Time Allowed"  # no table: English labels
    assert for_language("English").duration(60) == "1 hour" and for_language("English").duration(90) == "1 hour 30 minutes"


def test_a_word_made_of_two_fonts_is_shaped_font_by_font():
    """ReportLab shapes a whole word with its first font; an Indic word followed by a comma used to lose glyphs."""
    indic.ensure_font("devanagari")
    pdfmetrics.getFont("Serif") if "Serif" in pdfmetrics.getRegisteredFontNames() else None
    from app.export import pdf as pdfmod
    pdfmod._register_fonts()
    style = ParagraphStyle("t", fontName="Serif", fontSize=14, leading=18, shaping=1)
    word = indic.markup("क्षेत्रफल")
    plain = RLParagraph(word, style)
    plain.wrap(500, 100)
    joined = RLParagraph(word + ",", style)
    joined.wrap(500, 100)
    comma = pdfmetrics.stringWidth(",", "Serif", 14)
    assert abs((joined.getActualLineWidths0()[0] - plain.getActualLineWidths0()[0]) - comma) < 0.6


def test_indic_markup_wraps_runs_in_script_fonts_and_leaves_other_text_alone():
    out = indic.markup("(a) " + HINDI_PROSE + " x^2")
    assert out.count("<font name=") == 1 and out.startswith("(a) <font name=\"Indic-devanagari\">")
    assert out.endswith("</font> x^2")  # the danda is inside the run, the Latin tail outside
    assert 'Indic-devanagari-Bold"' in indic.markup(HINDI_PROSE, bold=True)
    assert indic.markup("Only English here") == "Only English here"
    assert indic.markup("हिन्दी, తెలుగు").count("<font") == 2  # two scripts, two fonts


def test_math_with_indic_words_falls_back_to_plain_text_instead_of_missing_glyphs():
    assert render_math("x^2 + 1") is not None
    assert render_math(r"\text{क्षेत्रफल} = 5") is None


# ----------------------------------------------------------------------------- DOCX

def test_docx_marks_runs_for_complex_script_font_size_bold_and_language():
    doc = Document(io.BytesIO(build_docx(_paper("Telugu", TELUGU_PROSE), "combined")))
    styles = doc.styles["Normal"].element.xml
    assert 'w:cs="Nirmala UI"' in styles and "<w:szCs" in styles
    runs = [r for p in doc.paragraphs for r in p.runs if languages.has_indic(r.text)]
    assert runs
    order = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "color", "position", "sz", "szCs", "lang"]
    for r in runs:
        rpr = r._r.rPr
        tags = [c.tag.rsplit("}", 1)[1] for c in rpr]
        assert "rFonts" in tags and "lang" in tags
        assert tags == sorted(tags, key=order.index), tags  # Word wants run properties in schema order
        if "sz" in tags:
            sz = next(c for c in rpr if c.tag.endswith("}sz"))
            szcs = next(c for c in rpr if c.tag.endswith("}szCs"))
            assert sz.attrib == szcs.attrib
    bold = [r for r in runs if r.bold]
    assert bold and all(r._r.rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bCs") is not None
                        for r in bold)
    xml = doc.element.xml
    assert 'w:bidi="te-IN"' in xml


def test_docx_text_and_labels_are_in_the_papers_language():
    doc = Document(io.BytesIO(build_docx(_paper("Hindi"), "combined")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert HINDI_PROSE in text and "समय" in text and "पूर्णांक" in text and "खंड A" in text
    assert "Time Allowed" not in text and "SECTION" not in text
    footer = "".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "पृष्ठ" in footer and "Page" not in footer


# ----------------------------------------------------------------------------- writing

def test_a_question_in_english_is_rejected_when_the_paper_is_in_an_indian_language():
    q = Question(id="q1", type="very_short", text="Define photosynthesis and explain why plants need chlorophyll.",
                 marks=2, answer=Answer(text="Plants make food."))
    assert "Hindi" in language_problem(q, "Hindi")
    assert language_problem(q, "English") is None and language_problem(q, "French") is None
    ok = Question(id="q1", type="very_short", text=HINDI_PROSE, marks=2, answer=Answer(text=HINDI_PROSE))
    assert language_problem(ok, "Hindi") is None
    maths = Question(id="q2", type="numerical", text="$x^2 - 5x + 6 = 0$", marks=2, answer=Answer(text="$x=2,3$"))
    assert language_problem(maths, "Hindi") is None  # too few letters to judge


def test_the_prompt_asks_for_the_papers_script_and_latin_diagram_labels():
    rule = language_rule("Telugu")
    assert "Telugu" in rule and "తెలుగు" in rule and "telugu" in rule and "Latin letters" in rule
    assert language_rule("English") == "Write in English."
    assert "French" in language_rule("French")


_NOUNS = ["पत्ती", "जड़", "तना", "फूल", "बीज", "फल", "मिट्टी", "पानी", "हवा", "धूप", "कोशिका", "ऊतक", "अंग", "रक्त",
          "हृदय", "फेफड़े", "गुर्दे", "मस्तिष्क", "हड्डी", "पेशी"]


class _Model(LLMProvider):
    """Answers like a model asked for a Hindi paper: slots in Hindi, translations in Hindi."""

    name, model = "hindi-test", "m-1"

    def __init__(self, english_first=False):
        super().__init__(2)
        self.english_first = english_first
        self.flags = {"spoiled": False}  # shared with the per-school copy the app makes of this provider
        self.calls = []

    def _complete(self, system, user, images, json_mode, max_tokens, temperature):
        self.calls.append(system[:40])
        if "translate" in system.lower():
            return json.dumps({"texts": ["अनुवादित: " + t for t in json.loads(user)]}, ensure_ascii=False)
        m = re.search(r"SLOTS TO WRITE:\n(\[.*\])", user)
        out = []
        for s in json.loads(m.group(1)):
            n = s["slot"]
            a, b, c = (_NOUNS[(n * 3 + k) % len(_NOUNS)] for k in range(3))
            text = f"{a} और {b} के बीच {c} का संबंध समझाइए ({n + 1})।"
            if self.english_first and n == 0 and not self.flags["spoiled"]:
                self.flags["spoiled"] = True  # only the first attempt comes back in English
                text = "Describe the process number three in plants and give reasons."
            item = {"slot": n, "text": text, "answer": {"text": "उत्तर: " + text}}
            if s["type"] == "mcq":
                item["options"] = [f"विकल्प {n}-{k} " + "अ" * (k + 1) for k in range(4)]
                item["answer"] = {"correct_option": 1, "text": "दूसरा विकल्प सही है।"}
            else:
                item["answer"]["steps"] = [{"text": f"चरण {i + 1} अंक {n}", "marks": s["marks"] / 2} for i in range(2)]
            out.append(item)
        return json.dumps({"questions": out}, ensure_ascii=False)


def _wait(client, url, headers, timeout=60):
    end = time.time() + timeout
    while time.time() < end:
        data = client.get(url, headers=headers).json()
        if data["status"] in ("ready", "failed"):
            return data
        time.sleep(0.2)
    raise AssertionError("timeout")


def test_a_hindi_paper_from_a_model_is_translated_banked_by_language_and_exported(settings):
    app = create_app(settings)
    with TestClient(app) as c:
        h = register(c)
        state = app.state.ctx
        state.provider = _Model()
        body = {"subject": "Biology", "pattern_id": "unit-test-20", "seed": 4, "language": "Hindi",
                "syllabus_text": "Photosynthesis\nRespiration"}
        plan = c.post("/api/papers/blueprint", headers=h, json=body).json()["writers"]
        assert plan["code"] == 0 and plan["language"]["needs_ai"] and not plan["language"]["will_be_english"]
        created = c.post("/api/papers", headers=h, json=body)
        paper = _wait(c, f"/api/papers/{created.json()['id']}", h)
        assert paper["status"] == "ready", paper["message"]
        assert paper["quality"]["failed"] == [] and paper["quality"]["writers"] == {"model": 11}
        content = paper["content"]
        assert content["meta"]["language"] == "Hindi"
        assert content["meta"]["instructions"][0].startswith("अनुवादित:")  # the pattern's wording was translated
        assert content["sections"][0]["title"].startswith("अनुवादित:")
        assert languages.has_indic(content["sections"][0]["questions"][0]["text"])
        with state.db.session() as db:
            assert {r.language for r in db.query(BankQuestion).all()} == {"hi"}
        owner = c.get("/api/auth/me", headers=h).json()["id"]
        assert len(_bank_items(state, owner, "Biology", "Hindi")) == 11
        assert _bank_items(state, owner, "Biology", "English") == []  # an English paper never reuses Hindi questions
        assert _bank_items(state, owner, "Biology", "Telugu") == []
        for fmt in ("pdf", "docx"):
            r = c.get(f"/api/papers/{paper['id']}/export", headers=h, params={"format": fmt, "mode": "combined"})
            assert r.status_code == 200
            assert "filename*=UTF-8''" in r.headers["content-disposition"] and r.headers["content-disposition"].isascii()
        z = c.get(f"/api/papers/{paper['id']}/export-all", headers=h, params={"sets": 1})
        assert z.status_code == 200 and len(zipfile.ZipFile(io.BytesIO(z.content)).namelist()) == 4


def test_english_answers_to_a_hindi_request_are_sent_back_for_rewriting(settings):
    app = create_app(settings)
    with TestClient(app) as c:
        h = register(c)
        model = _Model(english_first=True)
        app.state.ctx.provider = model
        body = {"subject": "Biology", "pattern_id": "unit-test-20", "seed": 4, "language": "Hindi",
                "syllabus_text": "Photosynthesis"}
        paper = _wait(c, f"/api/papers/{c.post('/api/papers', headers=h, json=body).json()['id']}", h)
        assert paper["status"] == "ready" and model.flags["spoiled"]
        assert paper["quality"]["failed"] == []
        assert all(languages.has_indic(q["text"]) for s in paper["content"]["sections"] for q in s["questions"])


def test_without_ai_a_hindi_paper_is_in_english_and_says_so(client):
    h = register(client)
    body = {"subject": "Mathematics", "pattern_id": "unit-test-20", "seed": 4, "language": "Hindi"}
    plan = client.post("/api/papers/blueprint", headers=h, json=body).json()["writers"]
    assert plan["language"]["will_be_english"]
    paper = _wait(client, f"/api/papers/{client.post('/api/papers', headers=h, json=body).json()['id']}", h)
    assert paper["status"] == "ready"
    assert any("cannot write Hindi" in w for w in paper["quality"]["warnings"]) and paper["quality"]["status"] == "review"


def test_templates_and_the_english_bank_are_not_used_for_other_languages():
    engine = HybridEngine(LLMEngine(_Model()), OfflineEngine(1))
    from app.exam.llm_engine import GenContext, SlotJob
    from app.exam.blueprint import build_blueprint
    pattern = get_preset("unit-test-20")
    for language, expected in (("English", True), ("Hindi", False)):
        req = GenerationRequest(subject="Mathematics", language=language, syllabus_text="Quadratic equations")
        slots = build_blueprint(req, pattern, resolve_topics(req, []))
        ctx = GenContext(req, pattern)
        routed = [engine.routes_to_code(ctx, SlotJob(s, [])) for s in slots]
        assert any(routed) == expected, language


def test_pattern_text_is_translated_once_and_text_already_in_the_language_is_left_alone():
    pattern = get_preset("unit-test-20")
    req = GenerationRequest(subject="Biology", language="Hindi", exam_name="अर्धवार्षिक परीक्षा")
    model = _Model()
    gen = PaperGenerator(LLMEngine(model), req, pattern, meter=None)
    names = gen._localised(list(pattern.instructions), pattern)
    assert all(t.startswith("अनुवादित:") for t in names["instructions"])
    assert names["titles"]["A"].startswith("अनुवादित:") and len(model.calls) == 1
    english = PaperGenerator(LLMEngine(model), req.model_copy(update={"language": "English"}), pattern)
    assert english._localised(list(pattern.instructions), pattern)["instructions"] == list(pattern.instructions)
    assert len(model.calls) == 1  # no call for English


# ----------------------------------------------------------------------------- documents and figures

def test_uploaded_hindi_material_is_recognised_and_offered_as_the_paper_language(client):
    h = register(client)
    text = ("# अध्याय 1\n\n" + (HINDI_PROSE + " ") * 12).encode()
    doc = client.post("/api/documents", headers=h, files=[("files", ("hindi.txt", text, "text/plain"))]).json()["documents"][0]
    end = time.time() + 30
    while time.time() < end:
        detail = client.get(f"/api/documents/{doc['id']}", headers=h).json()
        if detail["status"] in ("ready", "failed"):
            break
        time.sleep(0.2)
    assert detail["status"] == "ready" and detail["stats"]["script"] == "Devanagari" and detail["stats"]["language_hint"] == "Hindi"
    meta = client.get("/api/meta", headers=h).json()
    assert {"name": "Hindi", "native": "हिन्दी", "code": "hi", "script": "devanagari"} in meta["languages"]


def test_indic_letters_are_left_out_of_figures_rather_than_drawn_as_boxes():
    from app.diagrams import render_png
    from app.diagrams.render import _latin_only
    spec = {"type": "chart", "kind": "bar", "categories": ["क", "B"], "series": [{"name": "ऊँचाई", "values": [1, 2]}],
            "title": "ऊँचाई h", "caption": "चित्र 1"}
    cleaned = _latin_only(spec)
    assert cleaned["categories"] == ["", "B"] and cleaned["title"] == "h" and cleaned["caption"] == "चित्र 1"
    assert render_png(spec).data[:4] == b"\x89PNG"


def test_the_computer_answer_check_recognises_the_hindi_word_for_solve():
    from app.exam.verify import verify_question
    good = Question(id="q1", type="numerical", text="समीकरण $x^2 - 5x + 6 = 0$ को हल कीजिए।", marks=2,
                    answer=Answer(text="$x = 2$ या $x = 3$"))
    bad = good.model_copy(update={"answer": Answer(text="$x = 1$ या $x = 3$")})
    assert verify_question(good).status == "checked"
    assert verify_question(bad).status == "flagged"


def test_download_names_keep_the_indic_title_and_the_header_stays_ascii():
    from app.api.papers import _attachment, _filename
    name = _filename("विज्ञान - अर्धवार्षिक परीक्षा", "question", "A", "pdf")
    assert name == "विज्ञान_-_अर्धवार्षिक_परीक्षा_question_paper.pdf"
    header = _attachment(name)
    assert header.isascii() and header.startswith('attachment; filename="question_paper.pdf"') and "filename*=UTF-8''%E0%A4%B5" in header
    assert _filename("Maths: Half-Yearly", "combined", "B", "docx") == "Maths_Half-Yearly_teachers_copy_setB.docx"
    assert "\r" not in _attachment(_filename("x\r\nSet-Cookie: a=b", "question", "A", "pdf"))


# ----------------------------------------------------------------------------- Kannada subject, Kannada material

KN_NOTES = ("# ಅಧ್ಯಾಯ 1: ದ್ಯುತಿಸಂಶ್ಲೇಷಣೆ\n\n"
            + "ದ್ಯುತಿಸಂಶ್ಲೇಷಣೆ ಎಂದರೆ ಹಸಿರು ಸಸ್ಯಗಳು ಸೂರ್ಯನ ಬೆಳಕು, ನೀರು ಮತ್ತು ಕಾರ್ಬನ್ ಡೈಆಕ್ಸೈಡ್ ಬಳಸಿ ಆಹಾರ ತಯಾರಿಸುವ ಪ್ರಕ್ರಿಯೆ. "
            "ಹರಿತ್ತು ಎಲೆಗಳಲ್ಲಿರುವ ಹಸಿರು ವರ್ಣದ್ರವ್ಯವಾಗಿದ್ದು ಬೆಳಕಿನ ಶಕ್ತಿಯನ್ನು ಹೀರಿಕೊಳ್ಳುತ್ತದೆ. "
            "ಪತ್ರರಂಧ್ರಗಳು ಎಲೆಗಳ ಮೇಲ್ಮೈಯಲ್ಲಿರುವ ಸೂಕ್ಷ್ಮ ರಂಧ್ರಗಳು; ಅವುಗಳ ಮೂಲಕ ಅನಿಲಗಳ ವಿನಿಮಯ ನಡೆಯುತ್ತದೆ.\n" * 4
            + "\n# ಅಧ್ಯಾಯ 2: ಉಸಿರಾಟ\n\n"
            + "ಉಸಿರಾಟ ಎಂದರೆ ಜೀವಕೋಶಗಳಲ್ಲಿ ಆಹಾರದ ಅಣುಗಳಿಂದ ಶಕ್ತಿಯನ್ನು ಬಿಡುಗಡೆ ಮಾಡುವ ಪ್ರಕ್ರಿಯೆ. "
            "ಮೈಟೊಕಾಂಡ್ರಿಯಾವನ್ನು ಕೋಶದ ಶಕ್ತಿ ಕೇಂದ್ರ ಎಂದು ಕರೆಯುತ್ತಾರೆ. ಆಮ್ಲಜನಕ ಇಲ್ಲದೆ ನಡೆಯುವ ಉಸಿರಾಟವನ್ನು ಅವಾಯುಶ್ವಸನ ಎನ್ನುತ್ತಾರೆ.\n" * 4)
KN_NOUNS = ["ಎಲೆ", "ಬೇರು", "ಕಾಂಡ", "ಹೂವು", "ಬೀಜ", "ಹಣ್ಣು", "ಮಣ್ಣು", "ನೀರು", "ಗಾಳಿ", "ಬಿಸಿಲು", "ಕೋಶ", "ಅಂಗಾಂಶ", "ರಕ್ತ", "ಹೃದಯ",
            "ಶ್ವಾಸಕೋಶ", "ಮೂತ್ರಪಿಂಡ", "ಮೆದುಳು", "ಮೂಳೆ", "ಸ್ನಾಯು", "ಚರ್ಮ"]


class _KannadaModel(LLMProvider):
    name, model = "kannada-test", "m-1"

    def __init__(self):
        super().__init__(2)
        self.seen = []  # shared with the per-school copy the app makes

    def _complete(self, system, user, images, json_mode, max_tokens, temperature):
        if "translate" in system.lower():
            return json.dumps({"texts": ["ಅನುವಾದ: " + t for t in json.loads(user)]}, ensure_ascii=False)
        self.seen.append(user.split("SOURCE EXCERPTS", 1)[-1])
        out = []
        for s in json.loads(re.search(r"SLOTS TO WRITE:\n(\[.*\])", user).group(1)):
            n = s["slot"]
            a, b, c = (KN_NOUNS[(n * 3 + k) % len(KN_NOUNS)] for k in range(3))
            text = f"{a} ಮತ್ತು {b} ನಡುವಿನ {c} ಸಂಬಂಧವನ್ನು ವಿವರಿಸಿ ({n + 1})."
            item = {"slot": n, "text": text, "answer": {"text": "ಉತ್ತರ: " + text}}
            if s["type"] == "mcq":
                item["options"] = [f"{KN_NOUNS[(n * 4 + k) % 20]} ಆಯ್ಕೆ {k + 1}" for k in range(4)]
                item["answer"] = {"correct_option": 1, "text": "ಎರಡನೆಯ ಆಯ್ಕೆ ಸರಿಯಾಗಿದೆ."}
            else:
                item["answer"]["steps"] = [{"text": f"ಹಂತ {i + 1}", "marks": s["marks"] / 2} for i in range(2)]
            out.append(item)
        return json.dumps({"questions": out}, ensure_ascii=False)


def _upload_text(client, headers, name, text):
    doc = client.post("/api/documents", headers=headers,
                      files=[("files", (name, text.encode(), "text/plain"))]).json()["documents"][0]
    end = time.time() + 30
    while time.time() < end:
        detail = client.get(f"/api/documents/{doc['id']}", headers=headers).json()
        if detail["status"] in ("ready", "failed"):
            return detail
        time.sleep(0.2)
    raise AssertionError("document did not finish")


def test_indic_words_match_on_their_root_so_case_endings_do_not_hide_a_passage():
    assert tokenize("ಹರಿತ್ತು") == tokenize("ಹರಿತ್ತಿನ") and tokenize("ಪತ್ರರಂಧ್ರ") == tokenize("ಪತ್ರರಂಧ್ರಗಳು")
    assert tokenize("ಉಸಿರಾಟ") != tokenize("ಹರಿತ್ತು")
    assert tokenize("plants growing quickly") == ["plant", "grow", "quickly"]  # English stemming is untouched


def test_a_kannada_subject_paper_in_kannada_from_kannada_notes(settings):
    app = create_app(settings)
    with TestClient(app) as c:
        h = register(c)
        model = _KannadaModel()
        app.state.ctx.provider = model
        doc = _upload_text(c, h, "kannada_science.txt", KN_NOTES)
        assert doc["status"] == "ready" and doc["stats"]["language_hint"] == "Kannada"
        assert [ch["title"] for ch in doc["chapters"]] == ["ಅಧ್ಯಾಯ 1: ದ್ಯುತಿಸಂಶ್ಲೇಷಣೆ", "ಅಧ್ಯಾಯ 2: ಉಸಿರಾಟ"]
        body = {"subject": "ಕನ್ನಡ", "language": "Kannada", "pattern_id": "unit-test-20", "seed": 5,
                "document_ids": [doc["id"]], "institution": "ಸರ್ಕಾರಿ ಪ್ರೌಢಶಾಲೆ"}
        plan = c.post("/api/papers/blueprint", headers=h, json=body).json()
        assert plan["sources"]["unmatched_topics"] == [] and plan["writers"]["language"]["needs_ai"]
        paper = _wait(c, f"/api/papers/{c.post('/api/papers', headers=h, json=body).json()['id']}", h)
        assert paper["status"] == "ready", paper["message"]
        assert paper["quality"]["failed"] == [] and paper["quality"]["writers"] == {"model": 11}
        questions = [q for s in paper["content"]["sections"] for q in s["questions"]]
        assert all(languages.indic_share(q["text"]) == 1.0 for q in questions)
        assert any("ಹರಿತ್ತು" in seen or "ಉಸಿರಾಟ" in seen for seen in model.seen)  # the model was given the Kannada passages
        pdf = c.get(f"/api/papers/{paper['id']}/export", headers=h, params={"format": "pdf", "mode": "combined"})
        fonts = " ".join(f[3] for p in fitz.open(stream=pdf.content, filetype="pdf") for f in p.get_fonts())
        assert "NotoSansKannada" in fonts
        docx = Document(io.BytesIO(c.get(f"/api/papers/{paper['id']}/export", headers=h,
                                         params={"format": "docx", "mode": "question"}).content))
        assert "ಸರ್ಕಾರಿ ಪ್ರೌಢಶಾಲೆ" in "\n".join(p.text for p in docx.paragraphs)


def test_a_topic_typed_in_another_language_than_the_material_is_flagged(client):
    h = register(client)
    doc = _upload_text(client, h, "kannada_science.txt", KN_NOTES)
    body = {"subject": "Biology", "pattern_id": "unit-test-20", "document_ids": [doc["id"]], "language": "Kannada",
            "syllabus_text": "Photosynthesis\nಉಸಿರಾಟ"}
    plan = client.post("/api/papers/blueprint", headers=h, json=body).json()
    assert plan["sources"]["unmatched_topics"] == ["Photosynthesis"]  # the Kannada topic matches, the English one cannot
